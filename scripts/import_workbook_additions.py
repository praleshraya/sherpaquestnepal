"""Merge the new workbook rows and embedded photos into the existing catalog."""
import json
import re
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('/Users/praleshrayamajhi/Downloads/SHERPA QUEST PRODUCT (4).xlsx')
CATALOG = ROOT / 'data/products.json'
products = json.loads(CATALOG.read_text())
by_id = {item['id']: item for item in products}
book = openpyxl.load_workbook(SOURCE, data_only=True)
freshened = set()


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def add(name, category, brand, description, options='', product_id=None):
    product_id = product_id or re.sub(r'[^a-z0-9]+', '-', clean(name).lower()).strip('-')
    if product_id in by_id:
        by_id[product_id].update(name=clean(name), category=category, brand=brand,
                                 description=clean(description), options=clean(options))
        if product_id not in freshened:
            by_id[product_id]['gallery'] = []
            by_id[product_id]['image'] = ''
            freshened.add(product_id)
        return by_id[product_id]
    item = dict(id=product_id, name=clean(name), category=category, brand=brand,
                description=clean(description), options=clean(options), sourceImage='', image='')
    products.append(item)
    by_id[product_id] = item
    return item


def save_photo(image, item, variant):
    raw = image._data()
    digest = sha256(raw).hexdigest()
    seen = item.setdefault('_seen', set())
    if digest in seen:
        return
    seen.add(digest)
    folder = ROOT / 'assets/products' / item['id']
    folder.mkdir(parents=True, exist_ok=True)
    name = re.sub(r'[^a-z0-9]+', '-', variant.lower()).strip('-') or 'view'
    sequence = 1 + sum(x['variant'] == variant for x in item.get('gallery', []))
    destination = folder / f'{name}-{sequence:02d}.jpg'
    original = Image.open(BytesIO(raw))
    original.thumbnail((1200, 1200))
    rgba = original.convert('RGBA')
    photo = Image.new('RGB', rgba.size, 'white')
    photo.paste(rgba, mask=rgba.getchannel('A'))
    photo.save(destination, 'JPEG', quality=88, optimize=True)
    src = destination.relative_to(ROOT).as_posix()
    item.setdefault('gallery', []).append(dict(src=src, variant=variant))
    if not item['image']:
        item['image'] = src


guards = book['KNEE ELBOW ']
guard_rows = {
    3: add('KP101EP01 Knee & Elbow Guard', 'Knee & elbow guards', 'UNBRANDED', guards.cell(3, 2).value),
    4: add('KP16EP16 MX Knee & Elbow Guard', 'Knee & elbow guards', 'UNBRANDED', guards.cell(4, 2).value),
    5: add('WT01 Motorcycle Lumbar Support', 'Riding protection', 'UNBRANDED', guards.cell(5, 2).value),
}
for image in guards._images:
    row = image.anchor._from.row + 1
    if row in guard_rows:
        save_photo(image, guard_rows[row], 'Product view')

mounts = book['PHONE HOLDER']
mount_rows = {
    2: add('MH01 Traillock Pro Phone Mount', 'Phone holders', 'TRAILLOCK', mounts.cell(2, 2).value),
    4: add('MH02 Traillock X-Pro Phone & Camera Mount', 'Phone holders', 'TRAILLOCK', mounts.cell(4, 2).value),
}
for image in mounts._images:
    row = image.anchor._from.row + 1
    if row in mount_rows:
        save_photo(image, mount_rows[row], 'Product view')

osah = book['OSAH ']
osah_rows = {
    3: add('OSAH Edge Pro 40L Duffel Bag', 'Luggage & bags', 'OSAH', osah.cell(3, 3).value, 'Black'),
    4: add('OSAH Scout Tank Bag 6L', 'Luggage & bags', 'OSAH', osah.cell(4, 3).value, 'Black'),
    9: add('OSAH 6L ADV Crash Bar / Tail Bag', 'Luggage & bags', 'OSAH', osah.cell(9, 3).value, 'Black | Green'),
    11: add('OSAH Outrider Hydration Bag', 'Luggage & bags', 'OSAH', osah.cell(11, 3).value, 'Black'),
    12: add('OSAH Stretch Straps', 'Touring accessories', 'OSAH', osah.cell(12, 3).value),
    13: add('OSAH 10L Drypak', 'Luggage & bags', 'OSAH', osah.cell(13, 3).value, 'Blue | Olive'),
    14: add('OSAH 20L Drypak', 'Luggage & bags', 'OSAH', osah.cell(14, 3).value, 'Blue | Olive'),
    17: add('ILM Magnetic Tank Bag', 'Luggage & bags', 'ILM', osah.cell(17, 3).value, 'Black'),
    18: add('ILM Motorcycle Adventure Backpack', 'Luggage & bags', 'ILM', osah.cell(18, 3).value, 'Black'),
    19: add('ILM Balaclava FM02', 'Riding protection', 'ILM', osah.cell(19, 3).value, 'Black'),
    21: add('Rhinowalk Multifunctional Vest', 'Riding protection', 'RHINOWALK', osah.cell(21, 3).value, 'Black'),
}
osah_row_three = 0
osah_row_nine = 0
for image in osah._images:
    row = image.anchor._from.row + 1
    # Several photos are placed above their text row in the workbook.
    if row == 3:
        osah_row_three += 1
        if osah_row_three > 1: row = 4
    if row == 8: row = 9
    if row == 9:
        osah_row_nine += 1
    if row == 10:
        save_photo(image, osah_rows[9], 'Green')
    elif row in osah_rows:
        variant = 'Green' if row == 9 and osah_row_nine >= 7 else 'Black' if row in (9, 17, 18, 19, 21) else 'Product view'
        save_photo(image, osah_rows[row], variant)

mjw = book['Sheet14']
box_bag = add('MJW 6L Crash Bar Bag', 'Luggage & bags', 'MJW', mjw.cell(2, 2).value, 'Grey | Black')
for image in mjw._images:
    row = image.anchor._from.row + 1
    save_photo(image, box_bag, 'Black' if row == 3 else 'Grey')

for item in products:
    item.pop('_seen', None)
CATALOG.write_text(json.dumps(products, ensure_ascii=False, indent=2) + '\n')
print(f'{len(products)} catalog products; {len(products) - 55} additions from {SOURCE.name}')
